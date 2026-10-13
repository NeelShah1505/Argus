import csv
import io
import uuid
from decimal import Decimal
from typing import List, Optional, Tuple

from packages.canonical import (
    AccuracyTier,
    BoundingBox,
    EvidenceCoordinates,
    ExtractionMetadata,
    Page,
    Table,
    TableCell,
    TextElement,
)


def extract_csv(file_bytes: bytes, doc_id: str) -> Tuple[List[TextElement], List[Table]]:
    text = file_bytes.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    if not rows:
        return [], []

    table_id = f"tbl_{uuid.uuid4().hex[:8]}"
    cells: List[TableCell] = []
    text_elements: List[TextElement] = []

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            raw_val = val.strip()
            cell_id = f"c_{table_id}_{r_idx}_{c_idx}"
            evidence = EvidenceCoordinates(
                document_id=doc_id,
                page_number=1,
                bbox=BoundingBox(
                    page_number=1,
                    x0=float(c_idx * 100),
                    y0=float(r_idx * 30),
                    x1=float((c_idx + 1) * 100),
                    y1=float((r_idx + 1) * 30),
                ),
                raw_text=raw_val,
                render_image_key=f"renders/{doc_id}/p1.png",
            )
            cell = TableCell(
                id=cell_id,
                table_id=table_id,
                row_idx=r_idx,
                col_idx=c_idx,
                is_header=(r_idx == 0),
                raw_text=raw_val,
                tier=AccuracyTier.EXTRACTED,
                evidence=evidence,
                extraction=ExtractionMetadata(engines_used=["csv_native"], agreement_score=1.0),
            )
            cells.append(cell)

    tbl = Table(
        id=table_id,
        document_id=doc_id,
        page_number=1,
        title="CSV Extracted Table",
        rows_count=len(rows),
        cols_count=max(len(r) for r in rows) if rows else 0,
        cells=cells,
        bbox=BoundingBox(page_number=1, x0=0.0, y0=0.0, x1=1000.0, y1=1000.0),
        tier=AccuracyTier.EXTRACTED,
    )
    return text_elements, [tbl]


def extract_xlsx(file_bytes: bytes, doc_id: str) -> Tuple[List[TextElement], List[Table]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    tables: List[Table] = []
    text_elements: List[TextElement] = []

    for sheet_idx, sheetname in enumerate(wb.sheetnames, start=1):
        ws = wb[sheetname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        table_id = f"tbl_{uuid.uuid4().hex[:8]}"
        cells: List[TableCell] = []

        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                if val is None:
                    continue
                raw_val = str(val).strip()
                cell_id = f"c_{table_id}_{r_idx}_{c_idx}"
                evidence = EvidenceCoordinates(
                    document_id=doc_id,
                    page_number=sheet_idx,
                    bbox=BoundingBox(
                        page_number=sheet_idx,
                        x0=float(c_idx * 120),
                        y0=float(r_idx * 28),
                        x1=float((c_idx + 1) * 120),
                        y1=float((r_idx + 1) * 28),
                    ),
                    raw_text=raw_val,
                    render_image_key=f"renders/{doc_id}/p{sheet_idx}.png",
                )
                cell = TableCell(
                    id=cell_id,
                    table_id=table_id,
                    row_idx=r_idx,
                    col_idx=c_idx,
                    is_header=(r_idx == 0),
                    raw_text=raw_val,
                    tier=AccuracyTier.EXTRACTED,
                    evidence=evidence,
                    extraction=ExtractionMetadata(engines_used=["openpyxl_native"], agreement_score=1.0),
                )
                cells.append(cell)

        if cells:
            tbl = Table(
                id=table_id,
                document_id=doc_id,
                page_number=sheet_idx,
                title=f"Sheet: {sheetname}",
                rows_count=len(rows),
                cols_count=max(len(r) for r in rows) if rows else 0,
                cells=cells,
                bbox=BoundingBox(page_number=sheet_idx, x0=0.0, y0=0.0, x1=1000.0, y1=1000.0),
                tier=AccuracyTier.EXTRACTED,
            )
            tables.append(tbl)

    return text_elements, tables


def extract_docx(file_bytes: bytes, doc_id: str) -> Tuple[List[TextElement], List[Table]]:
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    text_elements: List[TextElement] = []
    tables: List[Table] = []

    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        if p_text:
            text_elements.append(
                TextElement(
                    id=f"txt_{doc_id}_{p_idx}",
                    document_id=doc_id,
                    page_number=1,
                    text=p_text,
                    bbox=BoundingBox(page_number=1, x0=50.0, y0=float(p_idx * 20), x1=900.0, y1=float((p_idx + 1) * 20)),
                )
            )

    for t_idx, t in enumerate(doc.tables):
        table_id = f"tbl_{uuid.uuid4().hex[:8]}"
        cells: List[TableCell] = []
        num_rows = len(t.rows)
        num_cols = len(t.columns) if t.rows else 0

        for r_idx, row in enumerate(t.rows):
            for c_idx, cell in enumerate(row.cells):
                raw_val = cell.text.strip()
                cell_id = f"c_{table_id}_{r_idx}_{c_idx}"
                evidence = EvidenceCoordinates(
                    document_id=doc_id,
                    page_number=1,
                    bbox=BoundingBox(
                        page_number=1,
                        x0=float(c_idx * 150),
                        y0=float(200 + r_idx * 25),
                        x1=float((c_idx + 1) * 150),
                        y1=float(200 + (r_idx + 1) * 25),
                    ),
                    raw_text=raw_val,
                    render_image_key=f"renders/{doc_id}/p1.png",
                )
                cell_model = TableCell(
                    id=cell_id,
                    table_id=table_id,
                    row_idx=r_idx,
                    col_idx=c_idx,
                    is_header=(r_idx == 0),
                    raw_text=raw_val,
                    tier=AccuracyTier.EXTRACTED,
                    evidence=evidence,
                    extraction=ExtractionMetadata(engines_used=["docx_native"], agreement_score=1.0),
                )
                cells.append(cell_model)

        tbl = Table(
            id=table_id,
            document_id=doc_id,
            page_number=1,
            title=f"DOCX Table {t_idx + 1}",
            rows_count=num_rows,
            cols_count=num_cols,
            cells=cells,
            bbox=BoundingBox(page_number=1, x0=50.0, y0=200.0, x1=950.0, y1=float(200 + num_rows * 25)),
            tier=AccuracyTier.EXTRACTED,
        )
        tables.append(tbl)

    return text_elements, tables


def extract_pdf(file_bytes: bytes, doc_id: str) -> Tuple[List[Page], List[TextElement], List[Table]]:
    """
    Extract digital-native PDF using PyMuPDF and pdfplumber.
    """
    import fitz  # PyMuPDF
    import pdfplumber

    pages: List[Page] = []
    text_elements: List[TextElement] = []
    tables: List[Table] = []

    pdf_doc = fitz.open(stream=file_bytes, filetype="pdf")

    with pdfplumber.open(io.BytesIO(file_bytes)) as plumber_pdf:
        for p_idx, page in enumerate(pdf_doc, start=1):
            w = page.rect.width
            h = page.rect.height
            pages.append(
                Page(
                    id=f"page_{doc_id}_{p_idx}",
                    document_id=doc_id,
                    page_number=p_idx,
                    render_key=f"renders/{doc_id}/p{p_idx}.png",
                    width=w,
                    height=h,
                )
            )

            # Extract text blocks
            blocks = page.get_text("blocks")
            for b_idx, b in enumerate(blocks):
                x0, y0, x1, y1, text, block_no, block_type = b
                raw_block_text = text.strip()
                if raw_block_text:
                    # Normalize coords to [0, 1000]
                    norm_bbox = BoundingBox(
                        page_number=p_idx,
                        x0=float((x0 / w) * 1000.0),
                        y0=float((y0 / h) * 1000.0),
                        x1=float((x1 / w) * 1000.0),
                        y1=float((y1 / h) * 1000.0),
                    )
                    text_elements.append(
                        TextElement(
                            id=f"txt_{doc_id}_{p_idx}_{b_idx}",
                            document_id=doc_id,
                            page_number=p_idx,
                            text=raw_block_text,
                            bbox=norm_bbox,
                        )
                    )

            # Extract tables using pdfplumber on matching page
            if p_idx - 1 < len(plumber_pdf.pages):
                plumber_page = plumber_pdf.pages[p_idx - 1]
                found_tables = plumber_page.extract_tables()
                for t_idx, raw_table in enumerate(found_tables):
                    if not raw_table:
                        continue
                    table_id = f"tbl_{uuid.uuid4().hex[:8]}"
                    cells: List[TableCell] = []
                    num_rows = len(raw_table)
                    num_cols = max(len(r) for r in raw_table) if raw_table else 0

                    for r_idx, row in enumerate(raw_table):
                        for c_idx, val in enumerate(row):
                            val_str = str(val).strip() if val is not None else ""
                            cell_id = f"c_{table_id}_{r_idx}_{c_idx}"
                            cell_bbox = BoundingBox(
                                page_number=p_idx,
                                x0=float((c_idx / max(num_cols, 1)) * 1000.0),
                                y0=float((r_idx / max(num_rows, 1)) * 1000.0),
                                x1=float(((c_idx + 1) / max(num_cols, 1)) * 1000.0),
                                y1=float(((r_idx + 1) / max(num_rows, 1)) * 1000.0),
                            )
                            evidence = EvidenceCoordinates(
                                document_id=doc_id,
                                page_number=p_idx,
                                bbox=cell_bbox,
                                raw_text=val_str,
                                render_image_key=f"renders/{doc_id}/p{p_idx}.png",
                            )
                            cell_model = TableCell(
                                id=cell_id,
                                table_id=table_id,
                                row_idx=r_idx,
                                col_idx=c_idx,
                                is_header=(r_idx == 0),
                                raw_text=val_str,
                                tier=AccuracyTier.EXTRACTED,
                                evidence=evidence,
                                extraction=ExtractionMetadata(engines_used=["pdfplumber"], agreement_score=1.0),
                            )
                            cells.append(cell_model)

                    tbl = Table(
                        id=table_id,
                        document_id=doc_id,
                        page_number=p_idx,
                        title=f"Page {p_idx} Table {t_idx + 1}",
                        rows_count=num_rows,
                        cols_count=num_cols,
                        cells=cells,
                        bbox=BoundingBox(page_number=p_idx, x0=50.0, y0=50.0, x1=950.0, y1=950.0),
                        tier=AccuracyTier.EXTRACTED,
                    )
                    tables.append(tbl)

    return pages, text_elements, tables
