import io
import csv
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Table as RLTable, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import docx
import openpyxl


def generate_test_pdf() -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Acme Corp — Consolidated Financial Statement", styles["Title"]))
    elements.append(Spacer(1, 14))
    elements.append(Paragraph("For the Fiscal Year Ended December 31, 2024 (in USD Millions)", styles["Normal"]))
    elements.append(Spacer(1, 14))

    table_data = [
        ["Line Item", "FY2023", "FY2024"],
        ["Total Revenue", "383,285", "391,036"],
        ["Cost of Sales", "(214,137)", "(213,972)"],
        ["Gross Profit", "169,148", "177,064"],
        ["Operating Expenses", "(54,847)", "(56,120)"],
        ["Operating Income", "114,301", "120,944"],
    ]
    t = RLTable(table_data, colWidths=[200, 100, 100])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.navy),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ]
        )
    )
    elements.append(t)
    doc.build(elements)
    return buf.getvalue()


def generate_test_docx() -> bytes:
    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_heading("Beta Holdings Q3 Earnings Report", level=1)
    doc.add_paragraph("Consolidated Statement of Operations and Comprehensive Income.")

    table = doc.add_table(rows=4, cols=3)
    data = [
        ["Metric", "2023", "2024"],
        ["Net Sales", "$120,000", "$145,000"],
        ["Net Income", "$22,000", "$28,500"],
        ["Cash and Equivalents", "$45,000", "$52,000"],
    ]
    for r_idx, row in enumerate(data):
        for c_idx, val in enumerate(row):
            table.cell(r_idx, c_idx).text = val

    doc.save(buf)
    return buf.getvalue()


def generate_test_xlsx() -> bytes:
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"

    data = [
        ["Asset Class", "FY2023", "FY2024"],
        ["Cash & Cash Equivalents", "29965", "29943"],
        ["Marketable Securities", "31590", "35210"],
        ["Accounts Receivable", "29508", "30120"],
        ["Total Current Assets", "91063", "95273"],
    ]
    for row in data:
        ws.append(row)

    wb.save(buf)
    return buf.getvalue()


def generate_test_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    data = [
        ["Segment", "Revenue_2023", "Revenue_2024"],
        ["North America", "162560", "168400"],
        ["Europe", "94294", "101000"],
        ["Greater China", "72559", "66950"],
        ["Rest of Asia Pacific", "29615", "31200"],
    ]
    for row in data:
        writer.writerow(row)
    return buf.getvalue().encode("utf-8")
